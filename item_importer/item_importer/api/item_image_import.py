import os
import csv
import zipfile
import tempfile

import frappe
from frappe.utils.file_manager import get_file_path, save_file
from openpyxl import load_workbook


IMAGE_FIELD_TYPES = [
    "Attach",
    "Attach Image",
    "Image",
]

IMAGE_FIELD_PRIORITY = [
    "image",
    "user_image",
    "profile_image",
    "profile_picture",
    "profile_pic",
    "photo",
    "picture",
    "attachment",
    "attach",
]


@frappe.whitelist()
def get_image_import_options(reference_doctype):
    """
    Used by JS.

    Purpose:
    1. Check whether selected Reference Doctype has image/attach field.
    2. Find which field image should be mapped to.

    This function does NOT create any field.
    It also does NOT return match_by options.
    """

    if not reference_doctype:
        return {
            "has_image_field": False,
            "matched_field": "",
            "image_fields": [],
        }
    meta = frappe.get_meta(reference_doctype)

    image_fields = []

    for df in meta.fields:
        if df.fieldtype in IMAGE_FIELD_TYPES:
            image_fields.append({
                "fieldname": df.fieldname,
                "label": df.label or df.fieldname,
                "fieldtype": df.fieldtype,
            })

    matched_field = get_image_field(image_fields)

    return {
        "has_image_field": bool(image_fields),
        "matched_field": matched_field,
        "image_fields": image_fields,
    }


@frappe.whitelist()
def attach_images_from_data_import(data_import_name):
    """
    Main button function.

    Works for any selected Reference Doctype.

    User manually enters:
    custom_image_match_by

    That field must already exist in selected reference_doctype.
    """

    data_import = frappe.get_doc("Data Import", data_import_name)

    reference_doctype = data_import.reference_doctype

    if not data_import.import_file:
        frappe.throw("Please attach Excel/CSV file")

    if not data_import.custom_image_zip_file:
        frappe.throw("Please attach Image ZIP file")

    if not data_import.custom_image_match_by:
        frappe.throw("Please enter Image Match By")

    meta = frappe.get_meta(reference_doctype)

    options = get_image_import_options(reference_doctype)

    if not options.get("has_image_field"):
        frappe.throw(f"{reference_doctype} does not have Image/Attach field")

    matched_field = data_import.custom_matched_field 

    if not matched_field:
        frappe.throw("Matched Field is missing")

    if not meta.has_field(matched_field):
        frappe.throw(
            f"Matched Field '{matched_field}' does not exist in {reference_doctype}"
        )

    match_by = str(data_import.custom_image_match_by).strip() 
    validate_match_by_field(reference_doctype, match_by)

    import_file_path = get_file_path(data_import.import_file)
    zip_file_path = get_file_path(data_import.custom_image_zip_file)

    rows = read_import_file(import_file_path)

    if not rows:
        frappe.throw("No rows found in import file")

    image_folder = extract_zip_file(zip_file_path)
    image_map = get_image_map(image_folder)

    total = 0
    success = 0
    failed = 0
    logs = []

    for row in rows:
        total += 1

        try:
            match_value = get_row_value(row, match_by)

            if not match_value:
                failed += 1
                logs.append(f"Row {total}: '{match_by}' value missing in import file")
                continue

            doc_name = get_document_name(reference_doctype, match_by, match_value)

            if not doc_name:
                failed += 1
                logs.append(
                    f"Row {total}: {reference_doctype} not found for {match_by} = {match_value}"
                )
                continue

            image_key = normalize(match_value)

            if image_key not in image_map:
                failed += 1
                logs.append(f"{doc_name}: Image not found for {match_value}")
                continue

            image_path = image_map[image_key]

            with open(image_path, "rb") as f:
                content = f.read()

            saved_file = save_file(
                fname=os.path.basename(image_path),
                content=content,
                dt=reference_doctype,
                dn=doc_name,
                is_private=0,
            )

            frappe.db.set_value(
                reference_doctype,
                doc_name,
                matched_field,
                saved_file.file_url,
                update_modified=True,
            )

            success += 1
            logs.append(
                f"Image attached successfully"
            )

        except Exception as e:
            failed += 1
            logs.append(f"Row {total}: Failed - {str(e)}")

    log_text = (
        f"Total Rows: {total}\n"
        f"Success: {success}\n"
        f"Failed: {failed}\n\n"
        + "\n".join(logs)
    )

    data_import.db_set("custom_matched_field", matched_field)
    data_import.db_set("custom_image_import_log", log_text)

    frappe.db.commit()

    return {
        "total": total,
        "success": success,
        "failed": failed,
        "logs": logs,
    }


def validate_match_by_field(reference_doctype, match_by):
    """
    Accept:
    - name
    - any existing fieldname from selected reference_doctype

    Does not create any field.
    """

    if match_by == "name":
        return

    meta = frappe.get_meta(reference_doctype)

    if not meta.has_field(match_by):
        frappe.throw(
            f"Match By field '{match_by}' does not exist in {reference_doctype}"
        )


def get_document_name(reference_doctype, match_by, match_value):
    if match_by == "name":
        if frappe.db.exists(reference_doctype, match_value):
            return match_value
        return None

    return frappe.db.get_value(
        reference_doctype,
        {
            match_by: match_value
        },
        "name",
    )


def get_image_field(image_fields):
    if not image_fields:
        return ""

    for priority_field in IMAGE_FIELD_PRIORITY:
        for image_field in image_fields:
            if image_field.get("fieldname") == priority_field:
                return image_field.get("fieldname")

    return image_fields[0].get("fieldname")


def get_row_value(row, fieldname):
    if fieldname in row:
        return row.get(fieldname)

    for key, value in row.items():
        if normalize_header(key) == normalize_header(fieldname):
            return value

    return ""


def normalize_header(value):
    return (
        str(value or "")
        .strip()
        .lower()
        .replace(" ", "_")
        .replace("-", "_")
    )


def read_import_file(file_path):
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".csv":
        return read_csv(file_path)

    if ext in [".xlsx", ".xlsm"]:
        return read_excel(file_path)

    frappe.throw("Only CSV or XLSX file is supported")


def read_csv(file_path):
    rows = []

    with open(file_path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)

        for row in reader:
            clean_row = {}

            for key, value in row.items():
                if key:
                    clean_row[key.strip()] = str(value).strip() if value else ""

            if any(clean_row.values()):
                rows.append(clean_row)

    return rows


def read_excel(file_path):
    workbook = load_workbook(file_path)
    sheet = workbook.active

    headers = []
    rows = []

    for cell in sheet[1]:
        headers.append(str(cell.value).strip() if cell.value else "")

    for row in sheet.iter_rows(min_row=2, values_only=True):
        row_data = {}

        for index, value in enumerate(row):
            if index < len(headers):
                fieldname = headers[index]

                if fieldname:
                    row_data[fieldname] = str(value).strip() if value else ""

        if any(row_data.values()):
            rows.append(row_data)

    return rows


def extract_zip_file(file_path):
    if not zipfile.is_zipfile(file_path):
        frappe.throw("Uploaded file is not a valid ZIP file")

    extract_folder = tempfile.mkdtemp()

    with zipfile.ZipFile(file_path, "r") as zip_ref:
        zip_ref.extractall(extract_folder)

    return extract_folder


def get_image_map(folder):
    image_map = {}

    allowed_extensions = [
        ".png",
        ".jpg",
        ".jpeg",
        ".webp",
    ]

    for root, dirs, files in os.walk(folder):
        for file_name in files:
            name_without_ext, ext = os.path.splitext(file_name)

            if ext.lower() not in allowed_extensions:
                continue

            key = normalize(name_without_ext)
            image_map[key] = os.path.join(root, file_name)

    return image_map


def normalize(value):
    value = str(value or "").strip().lower()

    for ext in [".png", ".jpg", ".jpeg", ".webp"]:
        if value.endswith(ext):
            value = value[:-len(ext)]

    return (
        value
        .replace(" ", "")
        .replace("-", "")
        .replace("_", "")
        .replace(".", "")
    )